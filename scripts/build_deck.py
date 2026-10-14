"""Builds docs/Mini_Project_Presentation.pptx from the department's presentation template.

The template (a deck in the same format: title, problem statement, literature survey tables, objectives,
architecture, proposed solution, results, conclusion, references) keeps its layouts, logos, colours and fonts;
this script replaces the text, tables and pictures slide by slide.

Usage: python -m scripts.build_deck "path/to/template.pptx"
Needs python-pptx. The figures come from docs/report_assets (scripts/report_diagram.py, scripts/capture_screens.py).
"""
import copy
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(ROOT, "docs", "report_assets")
OUT = os.path.join(ROOT, "docs", "Mini_Project_Presentation.pptx")
TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser("~"), "Downloads", "rppg semVII updated.pptx")
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"a": A}
FONT = "Times New Roman"
q = lambda tag: f"{{{A}}}{tag}"

TITLE = "Bank Transaction Analytics and Fraud Detection"
STUDENTS = ["Jay Dhake  (A – 07)", "Samarth Ghanekar  (A – 10)", "Saahir Joshi  (A – 16)"]
GUIDE = "Name of the Guide:- Prof. Mrunali Desai"

prs = Presentation(TEMPLATE)
slides = list(prs.slides)


# ------------------------------------------------------------------ helpers
def shape_named(slide, name):
    return next(s for s in slide.shapes if s.name == name)


def shape_with(slide, predicate):
    return next(s for s in slide.shapes if predicate(s))


def make_run(donor_run, text, bold=None, italic=None):
    run = copy.deepcopy(donor_run)
    rpr = run.find("a:rPr", NS)
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if italic is not None:
        rpr.set("i", "1" if italic else "0")
    run.find("a:t", NS).text = text
    return run


def rewrite(shape, specs):
    """specs: list of (template_paragraph_index, [text | (text, bold) | (text, bold, italic), ...]).
    Each new paragraph is a copy of the template paragraph, with its runs replaced."""
    body = shape.text_frame._txBody
    originals = body.findall("a:p", NS)
    any_run = next(r for p in originals for r in p.findall("a:r", NS) if (r.find("a:t", NS).text or "").strip())
    for p in originals:
        body.remove(p)
    for tidx, runs in specs:
        src = originals[tidx]
        para = copy.deepcopy(src)
        donors = para.findall("a:r", NS)
        donor = next((r for r in donors if (r.find("a:t", NS).text or "").strip()), any_run)
        for r in donors:
            para.remove(r)
        end = para.find("a:endParaRPr", NS)
        for item in runs:
            text, bold, italic = (item, None, None) if isinstance(item, str) else (tuple(item) + (None, None))[:3]
            run = make_run(donor, text, bold, italic)
            if end is not None:
                end.addprevious(run)
            else:
                para.append(run)
        body.append(para)


def set_first_run(shape, paragraph_index, text):
    para = shape.text_frame._txBody.findall("a:p", NS)[paragraph_index]
    run = para.find("a:r", NS)
    run.find("a:t", NS).text = text
    for extra in para.findall("a:r", NS)[1:]:
        para.remove(extra)


def set_cell(tc, text, donor_run):
    txbody = tc.find("a:txBody", NS)
    paras = txbody.findall("a:p", NS)
    for extra in paras[1:]:
        txbody.remove(extra)
    para = paras[0]
    runs = para.findall("a:r", NS)
    run = runs[0] if runs else copy.deepcopy(donor_run)
    for extra in runs[1:]:
        para.remove(extra)
    run.find("a:t", NS).text = text
    if not runs:
        end = para.find("a:endParaRPr", NS)
        (end.addprevious(run) if end is not None else para.append(run))


def edit_table(shape, data, col_widths_in=None, drop_col=None):
    """data: header row then data rows. Rows and columns are added or removed to fit."""
    tbl = shape.table._tbl
    if drop_col is not None:
        grid = tbl.find("a:tblGrid", NS)
        grid.remove(grid.findall("a:gridCol", NS)[drop_col])
        for tr in tbl.findall("a:tr", NS):
            tr.remove(tr.findall("a:tc", NS)[drop_col])
    trs = tbl.findall("a:tr", NS)
    while len(trs) > len(data):
        tbl.remove(trs.pop())
    while len(trs) < len(data):
        new = copy.deepcopy(trs[-1])
        trs[-1].addnext(new)
        trs.append(new)
    donor = next(r for tr in trs[1:] for r in tr.iter(q("r")))
    for tr, row in zip(trs, data):
        for tc, text in zip(tr.findall("a:tc", NS), row):
            set_cell(tc, text, donor)
    if col_widths_in:
        for gc, w in zip(tbl.find("a:tblGrid", NS).findall("a:gridCol", NS), col_widths_in):
            gc.set("w", str(int(w * 914400)))
    for tr in trs:  # let PowerPoint grow each row to its text
        tr.set("h", str(int(0.3 * 914400)))


def remove(shape, slide=None):
    rids = shape._element.xpath(".//a:blip/@r:embed")
    shape._element.getparent().remove(shape._element)
    for rid in rids:  # drop the picture from the file as well, not just from the slide
        if slide is not None:
            slide.part.drop_rel(rid)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ------------------------------------------------------------------ 1 title
s = slides[0]
box = shape_named(s, "Google Shape;68;p1")
set_first_run(box, 3, "Mini Project")
set_first_run(box, 5, TITLE)
for i, name in enumerate(STUDENTS):
    set_first_run(box, 7 + i, name)
set_first_run(box, 10, GUIDE)
notes(s, "Introduce the project, the team and the guide. One sentence: a big-data system that finds fraud among 6.36 million bank transactions and lets an analyst search, investigate and upload data.")

# ------------------------------------------------------------------ 2 problem statement
s = slides[1]
box = shape_named(s, "Google Shape;75;p2")
rewrite(box, [
    (0, ["The Core Problem"]),
    (1, ["Fraud is extremely ", ("rare", True), ": in the PaySim data only ", ("8,197 of 6,362,604", True), " transactions (", ("0.13%", True),
         ") are fraudulent. Existing approaches rely on fixed rules or single-machine notebooks that cannot search, investigate or score data at scale."]),
    (2, []),
    (3, [("Rules miss most fraud:", True), " The dataset's own rule flags only 16 of 8,197 frauds (recall 0.2%)."]),
    (4, [("Accuracy hides the problem:", True), " A model that always says \"legitimate\" is 99.87% accurate yet useless, so precision and recall are needed."]),
    (5, [("Scale and speed:", True), " A Spark scan to look up one account takes 1.4 s; an analyst needs the answer in milliseconds."]),
    (6, [("Analysis stops at the model:", True), " Flagged transactions must be searched, investigated as cases, and new files must be scored; most solutions offer none of this."]),
    (7, []),
    (8, []),
])
notes(s, "Fraud is about 1 in 776 transactions. Accuracy is meaningless; we judge with precision and recall. The gap is a complete system: scale, millisecond search, a case workflow and file upload.")

# ------------------------------------------------------------------ 3-5 literature survey
HEAD = ["Title/ Year", "Methodology", "Findings", "Limitation", "Algorithms Used"]
WIDTHS = [1.75, 1.9, 1.9, 1.6, 1.3]
LIT = [
    [
        ["Random Forests (Breiman, 2001)",
         "Introduced an ensemble of randomised decision trees trained on bootstrap samples; predictions are averaged.",
         "Accurate and robust to noise; gives variable importance. Basis of the fraud model.",
         "Hard to interpret as a whole; needs care with extreme class imbalance.",
         "Bagging, random feature selection, decision trees"],
        ["Apache Spark: A Unified Engine for Big Data Processing (Zaharia et al., 2016)",
         "Describes Spark's design: distributed datasets, DataFrames, SQL and libraries on one engine.",
         "Parallel in-memory processing and a query optimiser run batch, SQL and ML workloads on large data.",
         "Built for batch jobs; scanning for each single lookup is too slow for interactive use.",
         "DataFrame API, Spark SQL, Catalyst optimiser"],
        ["MLlib: Machine Learning in Apache Spark (Meng et al., 2016)",
         "Presents Spark's distributed ML library with pipelines, feature transformers and scalable learners.",
         "Provides Random Forest and Logistic Regression that train on millions of rows within pipelines.",
         "A trained model must be exported to be served outside Spark.",
         "ML pipelines, Random Forest, Logistic Regression"],
    ],
    [
        ["PaySim: A Financial Mobile Money Simulator for Fraud Detection (Lopez-Rojas et al., 2016)",
         "Agent-based simulator calibrated on aggregated logs of a real mobile-money service.",
         "Gives a labelled benchmark of 6.36 million transactions; fraud appears only in transfers and cash-outs.",
         "Simulated data; the fraud pattern is simpler than real fraud.",
         "Agent-based simulation"],
        ["SMOTE: Synthetic Minority Over-sampling Technique (Chawla et al., 2002)",
         "Creates synthetic minority-class examples by interpolating between neighbouring minority points.",
         "Improves detection of the minority class on imbalanced data compared with plain undersampling.",
         "Changes the data distribution and can create unrealistic samples; not used here.",
         "k-nearest neighbours, oversampling"],
        ["The Precision-Recall Plot Is More Informative than the ROC Plot (Saito and Rehmsmeier, 2015)",
         "Compared ROC and precision-recall plots on imbalanced binary classification tasks.",
         "Precision-recall plots show the true effect of false positives when positives are rare.",
         "Guidance on evaluation only; no new model.",
         "Precision-recall curve, ROC curve"],
    ],
    [
        ["DuckDB: an Embeddable Analytical Database (Raasveldt and Muhleisen, 2019)",
         "Describes an in-process analytical database with columnar, vectorised query execution.",
         "Runs analytical SQL without a server and reads Parquet files directly. Used for the search index.",
         "Single-writer design; not meant for many concurrent writers.",
         "Columnar storage, vectorised execution"],
        ["The Adaptive Radix Tree: ARTful Indexing for Main-Memory Databases (Leis et al., 2013)",
         "Proposes an adaptive radix tree index whose nodes adapt their size to the keys stored.",
         "Fast exact and prefix lookups with low memory use; behind account search and suggestions.",
         "Designed for main memory; ordering follows the key bytes.",
         "Adaptive radix tree, path compression"],
    ],
]
for slide, rows in zip(slides[2:5], LIT):
    table_shape = shape_with(slide, lambda sh: sh.has_table)
    edit_table(table_shape, [HEAD] + rows, WIDTHS)
    table_shape.left, table_shape.top = Inches(0.74), Inches(0.95)
notes(slides[2], "Foundations: Random Forests for the model, Spark for the batch pipeline, MLlib for training at scale.")
notes(slides[3], "The dataset (PaySim), how others handle imbalance (SMOTE), and why we use precision-recall metrics.")
notes(slides[4], "The two papers behind the search: DuckDB as the analytical engine and the adaptive radix tree as the index.")

# ------------------------------------------------------------------ 6 objectives
s = slides[5]
OBJ = {
    "Google Shape;102;p6": ("Process Big Data at Scale", [(2, ["Ingest and clean 6.36 million transactions into a partitioned Parquet lake with Spark."]),
                                                          (3, ["Run Spark SQL analyses of where, when and how fraud occurs."])]),
    "Google Shape;103;p6": ("Detect Rare Fraud Accurately", [(2, []), (3, ["Train and compare Logistic Regression and Random Forest with Spark MLlib."]),
                                                              (4, ["Judge them with precision, recall and AUC-PR, not accuracy."])]),
    "Google Shape;105;p6": ("Enable Instant Search", [(2, ["Find every transaction of an account in milliseconds with an index."]),
                                                       (3, ["Suggest matching account ids while the user types."])]),
    "Google Shape;104;p6": ("Support Investigation and Upload", [(2, []), (3, ["Manage fraud cases with create, read, update and delete."]),
                                                                  (4, ["Score uploaded files of millions of rows and explore them."])]),
}
for name, (heading, bullets) in OBJ.items():
    shp = shape_named(s, name)
    blank = [(1, [])]
    first_bullet_template = bullets[0][0]
    rewrite(shp, [(0, [heading])] + blank + [(t, r) for t, r in bullets])
    for run in shp.text_frame._txBody.find("a:p", NS).findall("a:r", NS):
        run.find("a:rPr", NS).set("sz", "1800")
notes(s, "Four objectives: process the data at scale, detect the rare fraud accurately, make search instant, and support investigation and upload.")

# ------------------------------------------------------------------ 7 system architecture
s = slides[6]
for pic in [sh for sh in s.shapes if sh.shape_type == 13]:
    remove(pic, s)  # the template's diagram and its full-slide overlay
height = Inches(4.35)
image = os.path.join(ASSETS, "fig_architecture.png")
from PIL import Image
with Image.open(image) as im:
    ratio = im.size[0] / im.size[1]
width = int(height * ratio)
s.shapes.add_picture(image, (prs.slide_width - width) // 2, Inches(0.95), width, height)
notes(s, "Top half is the offline batch pipeline run once with Spark. Bottom half is the always-on application: a REST API that reads the index, the model and the stored results, and reads and writes cases and uploads.")

# ------------------------------------------------------------------ 8 proposed solution
s = slides[7]
box = shape_named(s, "Google Shape;118;p8")
rewrite(box, [
    (0, [("Offline Batch Pipeline: ", True), "Spark ingests 6.36 million transactions into a partitioned Parquet lake, runs 7 SQL analyses and trains the models."]),
    (1, [("Fraud Model: ", True), "A 60-tree Random Forest scores transfers and cash-outs from nine features, with 99.3% precision and 99.75% recall on held-out data."]),
    (2, [("Millisecond Search: ", True), "An indexed table finds an account's transactions in about 2 to 7 ms, with suggestions while the user types."]),
    (3, [("Case Management: ", True), "Create, read, update and delete cases in MongoDB with a status workflow, notes and one open case per transaction."]),
    (4, [("Upload and Batch Scoring: ", True), "Files of millions of rows are streamed, validated and scored in about a minute, then indexed so they can be explored."]),
])
notes(s, "Five parts: the batch pipeline, the fraud model, millisecond search, case management, and upload with batch scoring.")

# ------------------------------------------------------------------ 9 results
s = slides[8]
cap1, cap2 = shape_named(s, "Google Shape;124;p9"), shape_named(s, "Google Shape;126;p9")
note = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("★"))
tables = [sh for sh in s.shapes if sh.has_table]
set_first_run(cap1, 0, "Table 1: Model Comparison on the Held-out Test Set (831,437 transactions, 2,407 frauds)")
set_first_run(cap2, 0, "Table 2: Time to Find All Transactions of One Account (6.36 million rows)")
edit_table(tables[0], [
    ["Model", "Precision", "Recall", "F1", "AUC-PR", "Frauds Caught", "False Alarms"],
    ["Logistic Regression (baseline)", "0.043", "0.909", "0.083", "0.549", "2,189 / 2,407", "48,466"],
    ["Random Forest (selected) ★", "0.993", "0.998", "0.995", "0.998", "2,401 / 2,407", "17"],
    ["Dataset rule (isFlaggedFraud)", "1.000", "0.002", "0.003", "–", "4 / 2,407", "0"],
], col_widths_in=[2.5, 1.0, 0.9, 0.8, 1.0, 1.7, 1.4])
edit_table(tables[1], [
    ["Method", "Median", "95th Percentile", "Faster than Spark"],
    ["Spark SQL scan of the data lake", "1,441 ms", "1,730 ms", "1×"],
    ["DuckDB on the Parquet files", "276 ms", "329 ms", "5×"],
    ["DuckDB table, no index", "28 ms", "35 ms", "51×"],
    ["DuckDB table with index (used) ★", "2.5 ms", "2.8 ms", "574×"],
], drop_col=4, col_widths_in=[3.4, 1.6, 1.9, 1.7])
cap2.top = Inches(2.5)
tables[1].top = Inches(2.85)
set_first_run(note, 0, "★ Unseen generated files: 8.0M rows scored in about 30 s and indexed in about 42 s, finding 92.1% of frauds at 51.8% precision; "
                       "the 7.0M-row unlabelled file was scored in about 35 s. 144 automated tests pass.")
notes(s, "Table 1: the Random Forest finds 2,401 of 2,407 frauds with 17 false alarms; the baseline raises 48,466. Table 2: the indexed lookup is 574 times faster than a Spark scan. On unseen generated files recall stays high but precision drops to 51.8%, an honest result on a different dataset.")

# ------------------------------------------------------------------ 10 conclusion
s = slides[9]
box = shape_named(s, "Google Shape;135;p11")
rewrite(box, [
    (0, ["A Spark pipeline and a ", ("60-tree Random Forest", True), " find 99.75% of frauds with 99.3% precision on unseen test data, far above the dataset's own rule (0.2% recall)."]),
    (1, ["An ", ("index", True), " makes account lookup 574 times faster than a Spark scan, so search and suggestions answer in milliseconds."]),
    (2, ["Cases (CRUD) and ", ("streaming upload", True), " turn the model into an investigation tool; 8-million-row files are scored and explorable in about a minute."]),
    (3, ["Limits: simulated data, Spark in local mode, no login. Future work: ", ("streaming scoring, a cluster, authentication", True), " and per-transaction explanations."]),
])
old_picture = next(sh for sh in s.shapes if sh.shape_type == 13)
x, y, w = old_picture.left, old_picture.top, old_picture.width
remove(old_picture, s)
shot = os.path.join(ASSETS, "fig_overview_top.png")
with Image.open(shot) as im:
    ratio = im.size[0] / im.size[1]
s.shapes.add_picture(shot, x, y + Inches(0.5), w, int(w / ratio))
notes(s, "Summarise: accurate model, instant search, a full case workflow and upload at scale. Be open about the limits: simulated data, local-mode Spark and no authentication.")

# ------------------------------------------------------------------ 11 references
s = slides[10]
box = shape_named(s, "Google Shape;141;p12")
REFS = [
    ("[1] L. Breiman, “Random Forests,” ", "Machine Learning", ", vol. 45, no. 1, pp. 5–32, 2001."),
    ("[2] M. Zaharia ", "et al.", ", “Apache Spark: A Unified Engine for Big Data Processing,” Commun. ACM, vol. 59, no. 11, 2016."),
    ("[3] X. Meng ", "et al.", ", “MLlib: Machine Learning in Apache Spark,” J. Mach. Learn. Res., vol. 17, no. 34, 2016."),
    ("[4] E. A. Lopez-Rojas, A. Elmir and S. Axelsson, “PaySim: A financial mobile money simulator for fraud detection,” ", "Proc. EMSS", ", 2016."),
    ("[5] N. V. Chawla ", "et al.", ", “SMOTE: Synthetic Minority Over-sampling Technique,” J. Artif. Intell. Res., vol. 16, 2002."),
    ("[6] T. Saito and M. Rehmsmeier, “The Precision-Recall Plot Is More Informative than the ROC Plot…,” ", "PLoS ONE", ", vol. 10, no. 3, 2015."),
    ("[7] M. Raasveldt and H. Muhleisen, “DuckDB: an Embeddable Analytical Database,” ", "Proc. ACM SIGMOD", ", 2019."),
    ("[8] V. Leis, A. Kemper and T. Neumann, “The Adaptive Radix Tree: ARTful Indexing for Main-Memory Databases,” ", "Proc. IEEE ICDE", ", 2013."),
    ("[9] Apache Spark Documentation, Spark SQL and MLlib guides. ", "spark.apache.org/docs/latest", ""),
    ("[10] Apache Parquet Documentation. ", "parquet.apache.org/docs", ""),
    ("[11] MongoDB Documentation, Indexes (TTL and partial). ", "mongodb.com/docs/manual/indexes", ""),
    ("[12] FastAPI Documentation. ", "fastapi.tiangolo.com", ""),
    ("[13] DuckDB Documentation. ", "duckdb.org/docs", ""),
]
specs = []
for a_, b_, c_ in REFS:
    runs = [a_, (b_, False, True)] + ([c_] if c_ else [])
    specs.append((3, runs))
rewrite(box, specs)
for para in box.text_frame._txBody.findall("a:p", NS):
    pPr = para.find("a:pPr", NS)
    pPr.set("algn", "just")
    for r in para.findall("a:r", NS):
        r.find("a:rPr", NS).set("sz", "1200")
box.left, box.top = Inches(0.45), Inches(0.95)
box.width = Inches(9.1)
notes(s, "Foundational papers for the model, the engine, the dataset, the evaluation and the index, plus the documentation of the tools used.")

# ------------------------------------------------------------------ Times New Roman everywhere
def force_font(element):
    for rpr in list(element.iter(q("rPr"))) + list(element.iter(q("endParaRPr"))) + list(element.iter(q("defRPr"))):
        for tag in ("latin", "ea", "cs", "sym"):
            node = rpr.find(f"a:{tag}", NS)
            if node is None:
                node = etree.SubElement(rpr, q(tag))
            node.set("typeface", FONT)
        # schema order: fill/ln ... latin, ea, cs, sym must follow the fills and precede hlinkClick
        for tag in ("latin", "ea", "cs", "sym"):
            node = rpr.find(f"a:{tag}", NS)
            rpr.remove(node)
            hl = rpr.find("a:hlinkClick", NS)
            (hl.addprevious(node) if hl is not None else rpr.append(node))


def force_everywhere(element):
    """The bullets, the master and the layouts use Arial for their own text styles; make those Times New Roman too."""
    for node in element.iter():
        if isinstance(node.tag, str) and node.tag.split("}")[-1] in ("latin", "ea", "cs", "sym", "buFont") and node.get("typeface"):
            node.set("typeface", FONT)


for slide in slides:
    force_font(slide._element)
    force_everywhere(slide._element)
force_everywhere(prs.slide_master._element)
for layout in prs.slide_layouts:
    force_everywhere(layout._element)
for part in prs.part.package.iter_parts():  # themes, the notes master, table styles and the presentation defaults
    if str(part.partname).endswith(".xml"):
        if hasattr(part, "_element"):
            force_everywhere(part._element)
        elif hasattr(part, "_blob") and b'typeface="Arial"' in part.blob:
            part._blob = part.blob.replace(b'typeface="Arial"', b'typeface="Times New Roman"')
prs.core_properties.title = TITLE
prs.core_properties.author = "Jay Dhake, Samarth Ghanekar, Saahir Joshi"
prs.save(OUT)
print("saved", OUT)
