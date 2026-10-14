"""Builds docs/TY_Mini_Project_Report.docx in the layout of the department's TY mini project report template.

Needs python-docx (requirements-dev.txt) and the images in docs/report_assets (scripts/capture_screens.py and
scripts/report_diagram.py create the figures). Usage: python -m scripts.build_report
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ASSETS = os.path.join(ROOT, "docs", "report_assets")
OUT = os.path.join(ROOT, "docs", "TY_Mini_Project_Report.docx")
FONT = "Times New Roman"
TEXT_WIDTH_CM = 16.0

TITLE = "Bank Transaction Analytics and Fraud Detection"
AUTHORS = ["Jay Dhake  (A \u2013 07)", "Samarth Ghanekar  (A \u2013 10)", "Saahir Joshi  (A \u2013 16)"]
SUPERVISOR = "Prof. Mrunali Desai"
YEAR = "2026-27"

import json
import math

PAGES_FILE = os.path.join(ASSETS, "_pages.json")
try:
    with open(PAGES_FILE) as f:
        PAGES = json.load(f)
except OSError:
    PAGES = {}
EST = {"page": 1, "pos": 0.0, "found": {}}
PAGE_CM = 25.3  # usable height of an A4 page with these margins


def advance(height, keep_whole=False):
    """Moves the layout cursor down; starts a new page when the content does not fit."""
    if keep_whole and EST["pos"] + height > PAGE_CM:
        EST["page"] += 1
        EST["pos"] = 0.0
    EST["pos"] += height
    while EST["pos"] > PAGE_CM:
        EST["pos"] -= PAGE_CM
        EST["page"] += 1


def est_lines(text, width_cm=16.0):
    return max(1, math.ceil(len(text.replace("**", "")) / (width_cm * 5.5)))


def new_page():
    EST["page"] += 1
    EST["pos"] = 0.0


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.different_first_page_header_footer = True


def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = None
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


normal = doc.styles["Normal"]
set_font(normal, 12)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for name, size, align in (("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER), ("Heading 2", 12, WD_ALIGN_PARAGRAPH.LEFT), ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT)):
    st = doc.styles[name]
    set_font(st, size, True)
    st.font.italic = False
    st.paragraph_format.alignment = align
    st.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
for name, indent in (("toc 1", 0), ("toc 2", 0.6)):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)  # Word's built-in index styles
    st.base_style = doc.styles["Normal"]
    set_font(st, 12, name == "toc 1")
    st.paragraph_format.left_indent = Cm(indent)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


# ------------------------------------------------------------------ helpers
def P(text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, after=None, highlight=False, indent=None):
    para = doc.add_paragraph()
    para.alignment = align
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    if indent:
        para.paragraph_format.left_indent = Cm(indent)
    add_runs(para, text, bold, italic, size, highlight)
    advance(est_lines(text) * 0.56 * ((size or 12) / 12) + 0.3)
    return para


def add_runs(para, text, bold=False, italic=False, size=None, highlight=False):
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        r = para.add_run(part)
        r.bold = bold or i % 2 == 1
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if highlight:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def bullets(items, style="List Bullet"):
    for item in items:
        para = doc.add_paragraph(style=style)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.space_after = Pt(4)
        add_runs(para, item)
        advance(est_lines(item, 14.8) * 0.56 + 0.15)


def numbered(items):
    bullets(items, "List Number")


def chapter(number, title):
    doc.add_page_break()
    new_page()
    EST["found"][f"{number} {title}"] = EST["page"]
    advance(2.4)
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(0)
    label.paragraph_format.keep_with_next = True
    r = label.add_run(f"Chapter {number}")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_heading(title, 1)


def section(text):
    EST["found"][text] = EST["page"] if EST["pos"] + 2.5 < PAGE_CM else EST["page"] + 1
    advance(1.6, keep_whole=False)
    doc.add_heading(text, 2)


def figure(filename, caption, width_cm=15.5, max_height_cm=20.5):
    path = os.path.join(ASSETS, filename)
    from PIL import Image
    with Image.open(path) as im:
        w, h = im.size
    width = width_cm
    if width * h / w > max_height_cm:
        width = max_height_cm * w / h
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.paragraph_format.space_after = Pt(4)
    para.add_run().add_picture(path, width=Cm(width))
    advance(width * h / w + 1.2, keep_whole=True)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(14)
    r = cap.add_run(caption)
    r.bold = True
    r.font.size = Pt(10)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(head, rows, widths, size=10.5, center=True, caption=None):
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_after = Pt(4)
        r = cap.add_run(caption)
        r.bold = True
        r.font.size = Pt(10)
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        pa = c.paragraphs[0]
        pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pa.paragraph_format.space_after = Pt(2)
        r = pa.add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        shade(c, "E7ECF1")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            pa = cells[i].paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            pa.paragraph_format.space_after = Pt(2)
            pa.paragraph_format.line_spacing = 1.0
            add_runs(pa, str(v), size=size)
    for r_index, row in enumerate(t.rows):
        cells_text = list(head) if r_index == 0 else [str(v) for v in rows[r_index - 1]]
        lines = max(math.ceil(len(c.replace("**", "")) / max(widths[i] * 5.8, 1)) for i, c in enumerate(cells_text))
        advance(lines * 0.46 + 0.2)
    advance(1.2 if caption else 0.6)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def footer_page_numbers():
    para = sec.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        run = para.add_run()
        run.font.name = FONT
        run.font.size = Pt(11)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


footer_page_numbers()

# ------------------------------------------------------------------ cover page
for filename, width, after in (("cover_1.png", 11.5, 22),):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(after)
    para.add_run().add_picture(os.path.join(ASSETS, filename), width=Cm(width))
P("DEPARTMENT OF COMPUTER ENGINEERING", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
P(TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=24, after=22)
P("TY Mini Project Report", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=8)
P("by", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=10)
for a in AUTHORS:
    P(a, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
P("Supervisor", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
P(SUPERVISOR, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, after=28)
para = doc.add_paragraph()
para.alignment = WD_ALIGN_PARAGRAPH.CENTER
para.paragraph_format.space_after = Pt(16)
para.add_run().add_picture(os.path.join(ASSETS, "cover_2.png"), width=Cm(4.2))
P("Department of Computer Engineering", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=8)
P("K.J. Somaiya Institute of Technology, Ayurvihar, Sion, Mumbai- 400022", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
P(YEAR, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

# ------------------------------------------------------------------ index (table of contents field)
doc.add_page_break()
new_page()
P("INDEX", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=12)

INDEX = [
    (1, "Abstract", 3),
    (1, "1 Introduction", 4),
    (2, "1.1 Introduction", 4), (2, "1.2 Motivation", 4), (2, "1.3 Problem Statement", 5), (2, "1.4 Objectives", 5),
    (1, "2 Literature Survey", 6),
    (2, "2.1 Survey of Existing Systems", 6), (2, "2.2 Findings from Literature Survey", 7),
    (1, "3 Proposed System", 9),
    (2, "3.1 Introduction", 9), (2, "3.2 Proposed System", 9), (2, "3.3 The Dataset", 10), (2, "3.4 Machine Learning Model and Technology", 12),
    (2, "3.5 System Workflow", 14), (2, "3.6 Codebase Structure", 16),
    (1, "4 Result", 18),
    (2, "4.1 Results", 18), (2, "4.2 Screenshots", 21),
    (1, "5 Conclusion", 31),
    (1, "References", 33),
]
fld_begin = lambda: ("begin", None)
first = True
KEYS = {"Abstract": "Abstract", "1 Introduction": "1 Introduction", "2 Literature Survey": "2 Literature Survey", "3 Proposed System": "3 Proposed System",
        "4 Result": "4 Result", "5 Conclusion": "5 Conclusion", "References": "References"}
for i, (level, text, page) in enumerate(INDEX):
    page = PAGES.get(text, page)
    para = doc.add_paragraph(style=f"toc {level}")
    para.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    if i == 0:
        for kind, instr in (("begin", None), (None, ' TOC \\o "1-2" \\h \\z \\u '), ("separate", None)):
            run = para.add_run()
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = instr
            run._r.append(el)
    para.add_run(f"{text}\t{page}")
    if i == len(INDEX) - 1:
        run = para.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), "end")
        run._r.append(el)
P("(In Word, right-click the index and choose Update Field to refresh the page numbers.)", italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

# ------------------------------------------------------------------ abstract
doc.add_page_break()
new_page()
EST["found"]["Abstract"] = EST["page"]
P("Abstract", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=14)
P("Bank Transaction Analytics and Fraud Detection is a big-data analytics system that finds fraudulent bank transactions among millions of records and lets an analyst "
  "investigate them. Fraud is extremely rare: in the PaySim dataset used here, 8,197 of 6,362,604 transactions (0.13%, about one in 776) are fraudulent, so the fraud cannot be found by "
  "inspection and a model that always answers \"legitimate\" would be 99.87% accurate yet useless. The system therefore measures precision and recall rather than accuracy.")
P("The solution has two halves. An offline batch pipeline built on Apache Spark ingests the data into a partitioned Parquet data lake, runs seven Spark SQL analyses, and trains and "
  "evaluates two classifiers with Spark MLlib. The selected Random Forest of 60 decision trees detects 99.75% of the frauds in a held-out test set with 99.3% precision (2,401 of 2,407 frauds caught, "
  "17 false alarms), against a precision of 4.3% for a class-weighted Logistic Regression and a recall of 0.17% for the dataset's own rule. The online half is a web application backed by a REST API. "
  "It provides a transaction explorer that finds every transaction of an account among 6.36 million rows in a few milliseconds using an index, with suggestions while typing; a case-management "
  "workflow with full create, read, update and delete operations stored in MongoDB; and a file upload that streams, validates and scores files of millions of rows, builds a search index for each upload, "
  "and lets the Overview, Explorer and Cases screens work on any uploaded file.")
P("Because real bank data is confidential, the system was also exercised with two generated files of 8.0 and 7.0 million rows, one with fraud labels and one without. On the labelled file, which "
  "the model had never seen, it found 92.1% of the frauds with 51.8% precision, and each file was scored in about 30 to 35 seconds and indexed in about 40 seconds. "
  "The behaviour of the system is verified by 144 automated tests.")
P("**Keywords:** Fraud detection, Big data, Apache Spark, Random Forest, Class imbalance, Parquet, Search index, REST API, MongoDB.", align=WD_ALIGN_PARAGRAPH.LEFT)

# ------------------------------------------------------------------ chapter 1
chapter(1, "Introduction")
section("1.1 Introduction")
P("Banks and mobile-money providers process millions of transactions every day: customers transfer money, withdraw cash, deposit cash and pay merchants. A very small fraction of these "
  "transactions are fraudulent, made by a criminal who has taken control of a customer's account rather than by the customer. Detecting them quickly matters, because stolen money that has been "
  "cashed out is very hard to recover. Manual review of millions of records is impossible, and simple rules miss most fraud.")
P("This project, Bank Transaction Analytics and Fraud Detection, addresses the problem with a complete big-data system. A batch pipeline processes the full dataset of 6.36 million "
  "transactions with Apache Spark, discovers where, when and how fraud occurs, and trains a machine-learning model that scores every transaction with a fraud probability. On top of the results sits "
  "a web application on which an analyst can see the findings, look up any account in milliseconds, open and manage investigation cases, and upload their own transaction files to have every row checked by the model.")
P("The data used is PaySim, a public benchmark of simulated mobile-money transactions generated from the logs of a real service and containing a ground-truth fraud label. In it, fraud follows a clear "
  "pattern, account takeover followed by cash-out: the money of a victim is transferred to a mule account and then withdrawn as cash. The system exploits the balance behaviour that this leaves behind, "
  "such as the sender's account being emptied and the receiver's recorded balance not changing although money arrived.")

section("1.2 Motivation")
P("The motivation for this project comes from three observations. First, fraud datasets are a textbook example of class imbalance. With one fraud in 776 transactions, accuracy is a meaningless measure, "
  "and an analyst needs a model that is judged by precision (how many alerts are real) and recall (how many frauds are found). Second, the data is large enough that the way it is stored and queried "
  "matters: a full scan of 6.36 million rows with a distributed engine takes more than a second, which is too slow for an interactive search, while an index brings the same lookup down to a few milliseconds.")
P("Third, a model on its own does not help an analyst. In practice, flagged transactions must be searched, reviewed, discussed and closed. This needs an application with a case workflow, "
  "and with the ability to score new batches of transactions that arrive as files. The project was therefore designed to cover the whole path from raw data to an investigated case, using "
  "technologies from the big-data ecosystem: Spark for batch processing and learning, Parquet for storage, a document database for operational data, and an embedded analytical database for fast lookups.")

section("1.3 Problem Statement")
P("Existing approaches to transaction fraud either rely on fixed rules that catch very few frauds, or are built as single-machine notebooks that cannot handle millions of rows, report accuracy "
  "that hides the rare class, and stop at a trained model without any way to search the data, investigate a flagged transaction or score new data. In the PaySim data, the provider's own rule flags only 16 of "
  "8,197 frauds, and common search over this data needs a full scan.")
P("Therefore, the problem is to design a scalable system that (a) processes millions of transactions with big-data tools, (b) detects fraud with high precision and recall despite extreme class imbalance, "
  "(c) lets users find any account's transactions in milliseconds, (d) supports a complete create, read, update and delete case workflow, and (e) scores uploaded files of millions of rows and makes them "
  "explorable, all through a documented and tested API.")

section("1.4 Objectives")
P("The key objectives of the project are as follows:", align=WD_ALIGN_PARAGRAPH.LEFT)
numbered([
    "To ingest and clean a dataset of 6.36 million transactions into a partitioned, columnar Parquet data lake using Apache Spark.",
    "To analyse the data with Spark SQL and discover where, when and how fraud occurs (transaction types, hours of the day, amounts, balance behaviour and repeat receivers).",
    "To train and compare classification models with Spark MLlib on the imbalanced data, and to evaluate them with precision, recall, F1 and the area under the precision-recall curve rather than accuracy.",
    "To serve the trained model without Spark, so that single transactions and files of millions of rows can be scored quickly by the web application.",
    "To provide millisecond search over all transactions by account, with suggestions while the user types, using an index.",
    "To implement fraud case management with full create, read, update and delete operations, a status workflow, notes and history, stored in a database.",
    "To allow a user to upload a CSV file of millions of rows, have every row validated and scored, and explore, analyse and open cases on the uploaded file.",
    "To expose everything through a validated REST API and to verify the system with automated tests.",
])

# ------------------------------------------------------------------ chapter 2
chapter(2, "Literature Survey")
section("2.1 Survey of Existing Systems")
P("A range of existing approaches address transaction fraud and big-data processing, but each has limitations with respect to the objectives of this project.")
P("**Rule-based fraud engines:** Banks traditionally flag transactions with fixed rules, for example \"flag any transfer above a threshold\". In the PaySim data, the provider's rule (isFlaggedFraud) "
  "flags only 16 of 8,197 fraudulent transactions, a recall of 0.2%. Rules are easy to explain but are quickly evaded and cannot learn the balance behaviour that characterises fraud.")
P("**Enterprise fraud platforms:** Commercial fraud-management products offer real-time scoring and case tools. They are proprietary, expensive and closed, so they cannot be used for academic work, "
  "and their internals cannot be inspected or reproduced.")
P("**Single-machine notebooks:** Many public solutions to fraud datasets use a notebook with pandas and scikit-learn. This is adequate for a few hundred thousand rows, but loads all data into one "
  "process's memory, reports accuracy on a rare class, and offers no search, case workflow or upload. The widely used credit-card fraud dataset from the Machine Learning Group of ULB, for instance, "
  "has 284,807 transactions of which 492 are fraudulent, which is small enough to hide scalability problems.")
P("**Distributed engines used for lookups:** A distributed batch engine such as Spark is excellent for processing and learning, but when it is used to answer an individual lookup it rescans the data. "
  "In our measurement, finding the transactions of one account among 6.36 million rows with Spark took a median of 1,441 ms, against 2.5 ms with an index.")
P("**Class-imbalance techniques:** Oversampling methods such as SMOTE create synthetic minority examples, and undersampling discards majority examples. They are useful but change the data distribution; "
  "this project instead uses class weights for the baseline model, deep trees for the forest, and evaluation with precision, recall and the precision-recall curve, which is recommended for imbalanced problems.")
P("**Public simulated datasets:** Real transactions are confidential. PaySim is a simulator calibrated on aggregated logs of a real mobile-money service, which provides a labelled benchmark of "
  "6.36 million transactions. It has become a common dataset for fraud research, with the limitation that its fraud follows a simplified pattern.")

section("2.2 Findings from Literature Survey")
table(["No.", "Title", "Authors", "Technology / Algorithm Used", "Results / Remarks"], [
    ["1", "Random Forests", "Breiman, 2001", "Ensemble of decision trees, bagging", "Showed that averaging many randomised trees gives accurate, robust classifiers; forms the core of the proposed fraud model."],
    ["2", "Apache Spark: A Unified Engine for Big Data Processing", "Zaharia et al., 2016", "Apache Spark, DataFrames, SQL", "Describes parallel in-memory processing of large datasets with a SQL optimiser; basis of the ingestion and analytics pipeline."],
    ["3", "MLlib: Machine Learning in Apache Spark", "Meng et al., 2016", "Distributed machine learning, pipelines", "Provides scalable Random Forest and Logistic Regression with pipelines; used to train and evaluate the models."],
    ["4", "PaySim: A Financial Mobile Money Simulator for Fraud Detection", "Lopez-Rojas, Elmir and Axelsson, 2016", "Agent-based simulation", "Source of the 6.36 million-row labelled dataset; fraud appears only in transfer and cash-out transactions."],
    ["5", "SMOTE: Synthetic Minority Over-sampling Technique", "Chawla et al., 2002", "Oversampling", "Addresses class imbalance by creating minority examples; motivates the project's choice of weights and suitable metrics instead."],
    ["6", "The Precision-Recall Plot Is More Informative than the ROC Plot on Imbalanced Datasets", "Saito and Rehmsmeier, 2015", "Evaluation metrics", "Justifies reporting precision, recall and area under the precision-recall curve for rare-class problems."],
    ["7", "DuckDB: an Embeddable Analytical Database", "Raasveldt and Muhleisen, 2019", "Columnar vectorised engine", "An in-process analytical database that reads Parquet directly; used for the search index."],
    ["8", "The Adaptive Radix Tree: ARTful Indexing for Main-Memory Databases", "Leis, Kemper and Neumann, 2013", "ART index", "A tree index supporting exact and prefix lookups; the index type behind the millisecond account search and suggestions."],
], widths=[1.1, 3.9, 2.7, 3.1, 5.2], size=10)

# ------------------------------------------------------------------ chapter 3
chapter(3, "Proposed System")
section("3.1 Introduction")
P("The proposed system, Bank Transaction Analytics and Fraud Detection, is a big-data system that finds fraudulent transactions and supports their investigation. It answers four questions: "
  "where, when and how does fraud occur in the data; can a model find it with high precision and recall; how can an analyst find the transactions of any account instantly; and how can the analyst "
  "investigate flagged transactions and check new files of transactions.")
P("The system is built as an offline batch pipeline and an online web application connected through files and databases. The pipeline uses Apache Spark for ingestion, SQL analytics and model training. "
  "The application uses a REST API, an indexed search store, a document database for operational data, and a web interface. The trained model is exported from Spark so that the application can score "
  "transactions without Spark.")

section("3.2 Proposed System")
P("The system provides the following functions:", align=WD_ALIGN_PARAGRAPH.LEFT)
numbered([
    "**Data analytics:** seven Spark SQL analyses of 6.36 million transactions, shown on an Overview page.",
    "**Fraud detection model:** a Random Forest trained with Spark MLlib that scores TRANSFER and CASH_OUT transactions.",
    "**Millisecond search (Explorer):** lookup by account with suggestions while typing, an account summary and a transaction detail view.",
    "**Case management:** create, read, update and delete cases, with a status workflow, priority, assignee, notes and history.",
    "**Upload and batch scoring:** streaming validation and scoring of files of millions of rows, with accuracy measurement when the file contains labels.",
    "**Dataset switching:** the Overview, Explorer and Cases screens can work on the main dataset or on any uploaded file.",
])
P("The application is deliberately split so that the heavy work is done once, offline, and the web application only reads prepared results, indexes and the exported model. "
  "This keeps every user action fast and lets the web server run without Spark or Java.")
figure("fig_architecture.png", "Fig 3.1 Architecture of the system: offline batch pipeline and online application", width_cm=15.5)

section("3.3 The Dataset")
P("**Main dataset (PaySim).** The main dataset is PaySim, published on Kaggle. It contains 6,362,604 simulated mobile-money transactions over 31 days (744 hourly steps). It is labelled: each "
  "row carries an isFraud flag, and the dataset also contains a simple built-in rule flag. The raw file (471 MB) is cleaned and converted into a 304 MB Parquet data lake partitioned by transaction type. "
  "No values are missing and there are no duplicate rows; 16 rows with an amount of zero were removed during ingestion.")
table(["Column", "Meaning"], [
    ["step", "Hour of the simulation (1 to 744, i.e. 31 days)"],
    ["type", "TRANSFER, CASH_OUT, PAYMENT, CASH_IN or DEBIT"],
    ["amount", "Amount of money moved"],
    ["nameOrig / nameDest", "Sender and receiver account ids (C for customers, M for merchants)"],
    ["oldbalanceOrg, newbalanceOrig", "Sender balance before and after"],
    ["oldbalanceDest, newbalanceDest", "Receiver balance before and after"],
    ["isFraud", "Ground-truth label: 1 for fraud, 0 for legitimate"],
    ["isFlaggedFraud", "Flag from the dataset's simple built-in rule"],
], widths=[5.5, 10.5], center=False, caption="Table 3.1 Columns of the transaction data")
P("**Where the fraud is.** Fraud occurs only in two of the five transaction types: 4,097 frauds in TRANSFER and 4,100 in CASH_OUT, none in PAYMENT, CASH_IN or DEBIT. The pattern is account takeover "
  "followed by cash-out: a victim's balance is transferred to a mule account, which then withdraws it. This leaves characteristic balance behaviour:")
table(["Behaviour", "Fraud", "Legitimate"], [
    ["Sender account emptied to zero", "97.7%", "42.7%"],
    ["Receiver's recorded balance stays 0 although money arrived", "49.7%", "0.06%"],
    ["Sender's balances do not reconcile with the amount", "0.55%", "90.6%"],
    ["Average amount", "1.47 million", "0.31 million"],
], widths=[8.5, 3.5, 4], caption="Table 3.2 Behaviour of fraudulent and legitimate transfers and cash-outs in the main dataset")
P("**Generated demo files.** Real bank data is confidential, so two additional files were generated with a seeded generator written for the project: one of 8,000,074 rows with an isFraud column and one of "
  "7,000,008 rows without it. The generator draws the transaction type from a fixed mix, the hour from a daily activity curve, the amount from a log-normal distribution per type, keeps balances consistent, adds a small share "
  "of genuine transactions that empty the sender's account, and injects account-takeover fraud (about 0.2% of rows) in which a victim's balance moves to a reused mule account and is cashed out. "
  "Only the shape of the data is borrowed from PaySim; no PaySim row is copied. The files are used to demonstrate scale and the upload workflow, and, because the model never saw them, they also serve as a test on unseen data.")

section("3.4 Machine Learning Model and Technology")
P("**Pipeline stages.** The batch pipeline has five steps. (1) The CSV is downloaded. (2) Spark reads it with an explicit schema, checks quality, adds derived columns (day, hour of day, "
  "the sender and receiver balance errors, and whether the sender was emptied) and writes Parquet partitioned by transaction type. (3) Spark SQL runs seven analyses whose results are stored in MongoDB "
  "and as JSON files. (4) Spark MLlib trains and evaluates the models. (5) A search index is built from the lake.")
P("**Features and models.** Only TRANSFER and CASH_OUT transactions (2.77 million rows) are modelled, because fraud occurs only there. Each is described by nine features: the transaction is a transfer, "
  "the amount, the four account balances, the two balance errors (balance after plus amount minus balance before, for sender and receiver) and the hour of the day. "
  "The data is split 70% for training (1,938,956 rows) and 30% for testing (831,437 rows, of which 2,407 are fraud). Two models are trained: a Logistic Regression with scaled features and class weights "
  "(each fraud counts 337 times a legitimate row) as the baseline, and a Random Forest of 60 trees of depth 12, which is the selected model. A transaction is flagged when the forest's average "
  "fraud probability exceeds 0.5.")
P("**Serving the model.** After training, the 60 trees are read from the files saved by Spark and evaluated with NumPy. Rows are scored in blocks of 25,000 on several threads, which reaches about "
  "426,000 rows per second. A test confirms that the scores equal Spark's on all 2,424 stored test transactions, with a maximum difference of 0.00005.")
table(["Layer", "Technology", "Purpose"], [
    ["Batch processing", "Apache Spark (PySpark, Spark SQL)", "Ingestion, cleaning and seven analyses over 6.36 million rows"],
    ["Machine learning", "Spark MLlib", "Training and evaluation of Logistic Regression and Random Forest"],
    ["Storage", "Parquet files", "Compressed columnar data lake, partitioned by type"],
    ["Search index", "DuckDB with ART indexes", "Millisecond lookup and prefix suggestions on account ids"],
    ["Operational database", "MongoDB", "Analysis results, fraud cases (with notes and history) and upload records"],
    ["API", "FastAPI, Pydantic", "REST endpoints with input validation and generated documentation"],
    ["Scoring and upload", "NumPy, PyArrow", "Model scoring outside Spark; streaming CSV validation and scoring"],
    ["Interface", "HTML, CSS, JavaScript, Chart.js", "Overview, Explorer, Cases and Upload screens"],
    ["Quality", "pytest", "144 automated tests"],
], widths=[3.4, 5.6, 7], center=False, caption="Table 3.3 Technologies used")
P("**Why an index.** A DuckDB table holds all 6.36 million transactions with an index on the sender, receiver and transaction id. The index is an Adaptive Radix Tree, which supports exact lookups "
  "and prefix range scans, so the same structure answers an account lookup and the suggestions shown while the user types. Section 4.1 compares it with scanning approaches.")

section("3.5 System Workflow")
P("The end-to-end workflow, from raw data to an investigated case, proceeds as follows.")
P("**Step 1: Data preparation (offline).** Running the pipeline downloads the data, creates the Parquet lake, runs the Spark SQL analyses, trains the model and builds the search index. "
  "Results are stored in MongoDB and in JSON files, so the Overview works even without the database.", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
P("**Step 2: Review the findings (Overview).** The user opens the Overview page, which shows the dataset summary, the fraud by transaction type, hour of day, amount band and balance behaviour, "
  "the accounts that repeatedly receive fraud, the comparison of models with the confusion matrix and feature importance, the flagged test transactions, and a form that scores a single transaction.")
P("**Step 3: Find an account (Explorer).** The user types an account id. After each pause in typing, the browser asks the API for accounts that start with those characters, and the API answers "
  "from the index. Choosing an account returns every transaction the account sent or received, each scored by the model, together with an account summary. Selecting a row shows the full detail.")
P("**Step 4: Investigate (Cases).** From a transaction the user opens a case. A case stores a snapshot of the transaction and moves through a workflow (new, investigating, confirmed fraud, false positive, "
  "closed). The user can change status, priority and assignee, add notes, search and filter cases, and delete a case. MongoDB enforces that only one open case exists per transaction, and cases expire after seven days of inactivity.")
P("**Step 5: Score a file (Upload).** The user uploads a CSV. The server reads it in 16 MB blocks, validates every row, scores the TRANSFER and CASH_OUT rows, counts the flagged rows, and, when the file has an "
  "isFraud column, computes precision, recall and the confusion matrix. Invalid rows are skipped and reported with their row numbers. A progress bar shows the status. The scored file can be downloaded.")
P("**Step 6: Explore the uploaded file.** When scoring finishes, the server builds a search index and an analysis for the file. The dataset picker then lets the Overview, Explorer and Cases screens work on the upload, "
  "including opening cases from its rows. Uploads are private to the browser and are deleted after 24 hours.")

section("3.6 Codebase Structure")
P("The project is organised into modules for the batch pipeline, the serving layer, the backend API, the web interface and tests.")
table(["File/Module", "Description"], [
    ["spark_jobs/ingest.py", "Reads the CSV with an explicit schema, checks data quality, adds derived columns and writes the partitioned Parquet data lake."],
    ["spark_jobs/analytics.py", "Runs the seven Spark SQL analyses and publishes each result to MongoDB and to a JSON file."],
    ["spark_jobs/train_model.py", "Trains and evaluates the Logistic Regression and Random Forest models, writes metrics, feature importance, the threshold curve and the flagged test transactions, and saves the model."],
    ["serving/build_search_index.py", "Builds the DuckDB table with a stable transaction id and indexes on the transaction id, sender and receiver."],
    ["backend/app/main.py", "Creates the FastAPI application, registers the routers, and defines the health, results, flagged-transactions and predict endpoints; also serves the web interface."],
    ["backend/app/search.py", "Account search, prefix suggestions, transaction and account detail, using the index with bound parameters."],
    ["backend/app/forest.py", "Loads the trained trees and scores transactions with NumPy in blocks on several threads."],
    ["backend/app/cases.py", "Create, read, update and delete operations for fraud cases, with the status workflow, limits and atomic updates."],
    ["backend/app/uploads.py", "Streaming upload: validation, scoring, accuracy metrics, per-upload index and analysis, download and delete."],
    ["backend/app/datasets.py", "Lists the datasets for the picker and provides the analysis of an uploaded file."],
    ["backend/app/db.py", "MongoDB connection, unique and expiry indexes, and the anonymous workspace check."],
    ["frontend/index.html, app.js", "The single-page interface, tab navigation, Overview charts and the single-transaction check."],
    ["frontend/explorer.js, cases.js, upload.js, datasets.js", "Explorer with suggestions, case management, upload and the dataset picker."],
    ["scripts/generate_dataset.py", "Generates the synthetic demo files (with or without the fraud label)."],
    ["tests/", "144 automated tests of scoring, search, cases, uploads, datasets and the generator."],
], widths=[5.6, 10.4], center=False, size=10.5, caption="Table 3.4 Codebase structure")

# ------------------------------------------------------------------ chapter 4
chapter(4, "Result")
section("4.1 Results")
P("The system was developed and tested using Python, Apache Spark, MongoDB, DuckDB and FastAPI. All components, including data processing, model training, search, case management, upload and "
  "visualisation, work correctly, and the behaviour is checked by 144 automated tests.")
P("**Model performance.** The models were evaluated on the 831,437 test transactions that were not used in training (2,407 of them fraudulent). The Random Forest finds 2,401 of the 2,407 frauds and raises "
  "17 false alarms. The baseline Logistic Regression finds more frauds than the dataset's own rule but raises 48,466 false alarms, and the rule finds only 4 of the test frauds.")
table(["Model", "Precision", "Recall", "F1", "AUC-PR", "Frauds caught", "False alarms"], [
    ["Logistic Regression (baseline)", "0.043", "0.909", "0.083", "0.549", "2,189 / 2,407", "48,466"],
    ["Random Forest (selected)", "0.993", "0.998", "0.995", "0.998", "2,401 / 2,407", "17"],
    ["Dataset rule (isFlaggedFraud)", "1.000", "0.002", "0.003", "-", "4 / 2,407", "0"],
], widths=[4.4, 1.8, 1.6, 1.4, 1.7, 2.7, 2.4], size=10, caption="Table 4.1 Comparison of models on the held-out test set")
P("The most important features of the Random Forest are the sender's balance after the transaction (40.0%), the sender's balance error (19.5%) and the amount (10.7%). "
  "The precision and recall are unchanged for thresholds between 0.2 and 0.7, so the default threshold of 0.5 is a stable choice.")
P("**Search speed.** The same lookup, finding every transaction sent by an account among 6,362,604 transactions, was timed four ways. The indexed lookup is 574 times faster than scanning with Spark; "
  "through the API the search, including scoring the rows, is answered in about 2 to 7 ms.")
table(["Method", "Median time", "95th percentile"], [
    ["Spark SQL scanning the data lake", "1,441 ms", "1,730 ms"],
    ["DuckDB reading the Parquet files directly", "276 ms", "329 ms"],
    ["DuckDB table without an index", "28 ms", "35 ms"],
    ["DuckDB table with the index (used by the system)", "2.5 ms", "2.8 ms"],
], widths=[8.5, 3.5, 4], caption="Table 4.2 Time to find all transactions of one account")
P("**Upload performance and accuracy on unseen data.** Two generated files that the model had never seen were uploaded through the web interface. Scoring runs in a stream, so memory does not grow with the file size.")
table(["File", "Rows", "Size", "Scoring", "Indexing", "Flagged", "Precision / Recall"], [
    ["With fraud labels", "8,000,074", "565 MB", "about 30 s", "about 42 s", "28,561", "51.8% / 92.1%"],
    ["Without fraud labels", "7,000,008", "481 MB", "about 35 s", "about 38 s", "25,014", "not measurable"],
], widths=[3.2, 2.2, 1.8, 2.1, 2.1, 1.9, 2.7], size=10, caption="Table 4.3 Upload of generated files")
P("On the labelled file, 4,520,189 of the rows (PAYMENT, CASH_IN and DEBIT) are not scored because fraud does not occur in those types. Of the 3,479,885 scored rows, the model flags 28,561, catches 14,807 of the "
  "16,072 frauds and raises 13,754 false alarms. Precision is lower than on the main dataset because the generated file is a different dataset, with genuine look-alike transactions, that the model was not trained on.")
P("**Case management.** Cases can be created, listed, filtered, edited, annotated and deleted. Opening a second case for the same transaction is rejected by the database, a closed case can only be reopened, and "
  "cases are private to each browser. These rules, the validation of inputs and the limits are covered by the automated tests.")

section("4.2 Screenshots")
figure("fig_overview_top.png", "Fig 4.1 Overview: dataset summary and fraud by transaction type")
figure("fig_model.png", "Fig 4.2 Model comparison, confusion matrix, feature importance and threshold curve")
figure("fig_suggestions.png", "Fig 4.3 Explorer: matching accounts appear while the account id is typed")
figure("fig_explorer.png", "Fig 4.4 Explorer: account summary, transactions and transaction detail with the model score")
figure("fig_cases.png", "Fig 4.5 Case management: list, status counts and the detail of a case being worked")
figure("fig_upload.png", "Fig 4.6 Upload: result of scoring an 8-million-row file with accuracy against its labels")

# ------------------------------------------------------------------ chapter 5
chapter(5, "Conclusion")
P("The Bank Transaction Analytics and Fraud Detection system demonstrates a complete big-data solution, from a raw file of 6.36 million transactions to an application in which an analyst can search, "
  "investigate and score data. Spark processed and analysed the data and trained a Random Forest that finds 99.75% of the frauds in a held-out test set with 99.3% precision, far better than the "
  "dataset's own rule (0.2% recall) and than a class-weighted Logistic Regression (4.3% precision). Evaluating with precision and recall, rather than accuracy, was essential for a problem where fraud is one in 776 transactions.")
P("Choosing the right tool for each job was equally important. Parquet reduced the storage by a third and allows partition pruning; an index made the account lookup 574 times faster than a Spark scan, which "
  "makes an interactive explorer possible; MongoDB suits the nested, short-lived case documents and enforces rules such as one open case per transaction; and exporting the model from Spark let the web "
  "server score transactions at about 426,000 rows per second without Spark. The upload pipeline processes files of millions of rows in a stream and makes them explorable in little more than a minute.")
P("The system has limitations that are stated openly. The data is simulated and the fraud pattern, account takeover followed by cash-out, is simpler than real fraud, so the reported accuracy is not a claim about real banks; "
  "on a separately generated file the model's precision fell to 51.8% while recall remained 92.1%. Spark was run in local mode on one machine. The model scores only transfers and cash-outs, uses a fixed threshold of 0.5, "
  "and the application uses anonymous per-browser workspaces instead of user accounts.")
P("**Future scope.** The work can be extended with streaming ingestion and real-time scoring, a Spark cluster, authentication and roles, per-transaction explanations of the model's decision, threshold selection based on "
  "the cost of false alarms, a graph view of mule-account networks, and training on real, anonymised data.")

# ------------------------------------------------------------------ references
doc.add_page_break()
new_page()
EST["found"]["References"] = EST["page"]
P("References", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=12)
refs = [
    "L. Breiman, \"Random Forests,\" Machine Learning, vol. 45, no. 1, pp. 5-32, 2001.",
    "M. Zaharia et al., \"Apache Spark: A Unified Engine for Big Data Processing,\" Communications of the ACM, vol. 59, no. 11, pp. 56-65, 2016.",
    "X. Meng et al., \"MLlib: Machine Learning in Apache Spark,\" Journal of Machine Learning Research, vol. 17, no. 34, pp. 1-7, 2016.",
    "E. A. Lopez-Rojas, A. Elmir and S. Axelsson, \"PaySim: A financial mobile money simulator for fraud detection,\" in Proc. 28th European Modeling and Simulation Symposium (EMSS), 2016. Dataset available at: https://www.kaggle.com/datasets/ealaxi/paysim1",
    "N. V. Chawla, K. W. Bowyer, L. O. Hall and W. P. Kegelmeyer, \"SMOTE: Synthetic Minority Over-sampling Technique,\" Journal of Artificial Intelligence Research, vol. 16, pp. 321-357, 2002.",
    "T. Saito and M. Rehmsmeier, \"The Precision-Recall Plot Is More Informative than the ROC Plot When Evaluating Binary Classifiers on Imbalanced Datasets,\" PLoS ONE, vol. 10, no. 3, 2015.",
    "M. Raasveldt and H. Muhleisen, \"DuckDB: an Embeddable Analytical Database,\" in Proc. ACM SIGMOD International Conference on Management of Data, 2019.",
    "V. Leis, A. Kemper and T. Neumann, \"The Adaptive Radix Tree: ARTful Indexing for Main-Memory Databases,\" in Proc. IEEE 29th International Conference on Data Engineering (ICDE), 2013.",
    "Apache Spark Documentation, Spark SQL and MLlib guides. Available at: https://spark.apache.org/docs/latest/",
    "Apache Parquet Documentation. Available at: https://parquet.apache.org/docs/",
    "MongoDB Documentation, Indexes: TTL and partial indexes. Available at: https://www.mongodb.com/docs/manual/indexes/",
    "FastAPI Documentation. Available at: https://fastapi.tiangolo.com/",
    "DuckDB Documentation. Available at: https://duckdb.org/docs/",
]
for i, r in enumerate(refs, 1):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para.paragraph_format.left_indent = Cm(0.9)
    para.paragraph_format.first_line_indent = Cm(-0.9)
    para.paragraph_format.space_after = Pt(6)
    para.add_run(f"{i}.\t{r}")

def all_times_new_roman(document):
    """Every style and the document defaults use Times New Roman; theme-font attributes would override it, so they are removed."""
    styles_el = document.styles.element
    for rfonts in styles_el.iter(qn("w:rFonts")):
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), FONT)
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        rpr = rpr_default.find(qn("w:rPr")) if rpr_default is not None else None
        if rpr is not None and rpr.find(qn("w:rFonts")) is None:
            rf = OxmlElement("w:rFonts")
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(attr), FONT)
            rpr.insert(0, rf)
    for style in document.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH or style.type == WD_STYLE_TYPE.CHARACTER:
            if style.font is not None:
                style.font.name = FONT


all_times_new_roman(doc)

# ask Word to refresh the index (page numbers) when the file is opened
settings = doc.settings.element
update = OxmlElement("w:updateFields")
update.set(qn("w:val"), "true")
# the schema fixes where updateFields goes: before the header shape defaults, footnote settings and compat block
anchor = next((c for c in settings if c.tag in {qn(t) for t in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids")}), None)
if anchor is not None:
    anchor.addprevious(update)
else:
    settings.append(update)

with open(PAGES_FILE, "w") as f:
    json.dump(EST["found"], f)
doc.save(OUT)
print("saved", OUT, "estimated pages:", EST["page"])
